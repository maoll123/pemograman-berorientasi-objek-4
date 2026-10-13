import docx
from docx.oxml import parse_xml
import json

doc = docx.Document('Laporan_Praktikum_PBO_Modul_3_AutoRent.docx')
print("Document Paragraphs:", len(doc.paragraphs))
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        runs_desc = []
        for r in p.runs:
            runs_desc.append({
                'text': r.text,
                'bold': r.bold,
                'italic': r.italic,
                'font': r.font.name,
                'size': str(r.font.size),
                'color': str(r.font.color.rgb) if r.font.color else None
            })
        print(f"P{i} [Align: {p.alignment}] [Style: {p.style.name}]:")
        print("  Runs:", runs_desc[:3])

print("\n" + "="*50 + "\nTables detail:")
for idx, t in enumerate(doc.tables):
    print(f"\n--- Table {idx+1} (Rows: {len(t.rows)}, Cols: {len(t.columns)}) ---")
    for r_idx, r in enumerate(t.rows[:3]):
        cells_data = []
        for c in r.cells:
            # check cell shading
            shd = c._tc.xpath('.//w:shd')
            fill = shd[0].attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill') if shd else None
            cells_data.append(f"{c.text[:25]} (fill={fill})")
        print(f"  Row {r_idx}: {cells_data}")
